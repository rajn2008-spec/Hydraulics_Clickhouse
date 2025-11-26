import pytest
import openpyxl
import json
from app.logic.diagnostics import run_diagnostics_basic

# -----------------------------
# Mapping Layer (canonical → human-readable)
# -----------------------------
OUTPUT_MAP = {
    "v_pipe": "Velocity Drill Pipe (ft/min)",
    "v_ann": "Velocity Annulus (ft/min)",
    "ECD": "Equivalent Circulating Density (lb/gal)",
    "Et": "Transmission Efficiency (%)",
    "Re_pipe": "Reynolds Number (Pipe)",
    "Re_ann": "Reynolds Number (Annulus)",
    "flow_string": "Flow Regime (Pipe)",
    "flow_annulus": "Flow Regime (Annulus)",
    "mu_e_pipe": "Effective Viscosity (Pipe, cP)",
    "mu_e_ann": "Effective Viscosity (Annulus, cP)",
    "n_string": "Flow Behavior Index (Pipe)",
    "k_string": "Consistency Index (Pipe)",
    "n_annulus": "Flow Behavior Index (Annulus)",
    "k_annulus": "Consistency Index (Annulus)",
}

# -----------------------------
# Input Reader
# -----------------------------
def read_inputs(sheet):
    inputs = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        param, value = row
        if param is None:
            continue
        if isinstance(value, str):
            try:
                value = float(value)
            except ValueError:
                pass
        inputs[param] = value
    return inputs

# -----------------------------
# Output Writer
# -----------------------------
def write_outputs(sheet, outputs):
    for idx, (canonical_key, human_label) in enumerate(OUTPUT_MAP.items(), start=2):
        sheet.cell(row=idx, column=1, value=canonical_key)
        sheet.cell(row=idx, column=2, value=human_label)

        value = outputs.get(canonical_key, 0.0)

        # Column 3 = raw value
        sheet.cell(row=idx, column=3, value=value)

        # Column 4 = JSON if needed
        if isinstance(value, (dict, list)):
            sheet.cell(row=idx, column=4, value=json.dumps(value))
        else:
            sheet.cell(row=idx, column=4, value=value)

# -----------------------------
# Main Test Function
# -----------------------------
@pytest.mark.parametrize("excel_file", ["hydraulics_template.xlsx"])
def test_run_diagnostics_excel(excel_file):
    wb = openpyxl.load_workbook(excel_file)
    inputs_sheet = wb["Inputs"]
    outputs_sheet = wb["Outputs"]

    inputs = read_inputs(inputs_sheet)
    outputs = run_diagnostics_basic(inputs).model_dump(exclude_none=True)

    write_outputs(outputs_sheet, outputs)
    wb.save("hydraulics_data_results.xlsx")

    # ✅ Canonical key check
    assert outputs.get("v_pipe") is not None, "v_pipe (velocity_drill_pipe) was not calculated"

    expected_keys = [
        "v_pipe", "v_ann", "n", "k", "mu", "Re", "Remax",
        "Vcrit", "Qcrit", "dp", "dp_L", "ECD", "rhoe"
    ]

    for key in expected_keys:
        assert outputs.get(key) is not None, f"{key} was not calculated"
